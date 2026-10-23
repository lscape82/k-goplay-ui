# -*- coding: utf-8 -*-
"""권역_데이터_산출근거 Excel → data/region-reports.json (지역별 자료 다운로드용).
재실행 가능. 숫자는 공신력 출처(소상공인365·TOPIS·철도통계·서울열린데이터) 확정값.
"""
import io, json, os, sys
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

XL = sys.argv[1] if len(sys.argv) > 1 else \
    r"E:/2. 광고플레이(주)/kgoplay_UIUX/2026광고주용_제안서만들기/★내부자료_권역별데이터분석/권역_데이터_산출근거_2025년12개월.xlsx"
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(os.path.dirname(__file__), "..", "data", "region-reports.json")

BUS_SHEET = {  # 버스 월별 시트 ↔ 권역
    "버스_월별_강남": "강남역·강남대로", "버스_월별_명동": "명동·을지로입구",
    "버스_월별_홍대": "홍대입구역", "버스_월별_성수": "성수·연무장길",
    "버스_월별_광화문": "광화문·종로",
}

def num(v):
    if v is None or v == "":
        return None
    try:
        return round(float(v), 2)
    except Exception:
        return str(v).strip()

def rows(ws):
    return list(ws.iter_rows(values_only=True))

wb = openpyxl.load_workbook(XL, data_only=True)
reports = {}

# 기준 권역(클린 5개) — 유동인구_월별 시트 기준. 다른 시트의 긴 라벨(교통량 'D-35 ...')은 여기로 정규화.
BASE = [str(r[0]).strip() for r in rows(wb["유동인구_월별"])[1:] if r and r[0]]

def resolve(name):
    raw = str(name).strip()
    for b in BASE:
        if raw == b or raw.startswith(b) or b in raw:
            return b
    return raw

# 플랫폼 지역 랭킹 키(areaName: '강남역/강남대로','광화문','홍대','성수','명동/을지로입구')와 매칭할 짧은 토큰
MATCH = {"강남역·강남대로": "강남", "명동·을지로입구": "명동", "홍대입구역": "홍대",
         "성수·연무장길": "성수", "광화문·종로": "광화문"}

def region(name):
    k = resolve(name)
    g = reports.setdefault(k, {"label": k})
    if "matchKey" not in g:
        g["matchKey"] = MATCH.get(k, k.split("·")[0].split("/")[0])
    return g

# 요약 — 지표표
for r in rows(wb["요약"])[1:]:
    if not r or not r[0]:
        continue
    g = region(r[0]); g.setdefault("summary", [])
    g["summary"].append({"지표": str(r[1]).strip() if r[1] else "", "단위": str(r[2]).strip() if r[2] else "",
                          "값": num(r[3]), "출처": str(r[4]).strip() if len(r) > 4 and r[4] else ""})

# 월별 시트 (유동·교통·지하철)
MON = [("유동인구_월별", "foot"), ("교통량_월별", "traffic"), ("지하철_월별", "subway")]
for sheet, key in MON:
    if sheet not in wb.sheetnames:
        continue
    for r in rows(wb[sheet])[1:]:
        if not r or not r[0]:
            continue
        g = region(r[0])
        g[key + "Monthly"] = [num(x) for x in r[1:13]]
        if len(r) > 13:
            g[key + "Avg"] = num(r[13])

# 버스 — 정류장 합계 월별 + 정류장 목록
for sheet, reg in BUS_SHEET.items():
    if sheet not in wb.sheetnames:
        continue
    g = region(reg); stops = []; tot = [0.0] * 12
    for r in rows(wb[sheet])[1:]:
        if not r or not r[1]:
            continue
        mv = [num(x) or 0 for x in r[2:14]]
        for i, v in enumerate(mv[:12]):
            tot[i] += (v or 0)
        stops.append({"ars": str(r[0]).strip() if r[0] else "", "name": str(r[1]).strip(),
                      "avg": round(sum(mv[:12]) / 12, 0) if mv else None})
    g["busMonthly"] = [round(x, 0) for x in tot]
    g["busAvg"] = round(sum(tot) / 12, 0)
    g["busStops"] = sorted(stops, key=lambda s: -(s["avg"] or 0))

# 연령·성별
if "연령·성별" in wb.sheetnames:
    hdr = rows(wb["연령·성별"])[0]
    for r in rows(wb["연령·성별"])[1:]:
        if not r or not r[0]:
            continue
        g = region(r[0])
        g["age"] = {str(hdr[i]): num(r[i]) for i in range(1, 7) if hdr[i]}
        g["gender"] = {"남": num(r[7]), "여": num(r[8])}

# 대상매체
if "대상매체" in wb.sheetnames:
    for r in rows(wb["대상매체"])[1:]:
        if not r or not r[0] or not r[1]:
            continue
        g = region(r[0]); g.setdefault("media", [])
        g["media"].append(str(r[1]).strip())

os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, "w", encoding="utf-8") as f:
    json.dump({"_source": "권역_데이터_산출근거_2025년12개월.xlsx (소상공인365·서울TOPIS·철도통계·서울열린데이터)",
               "regions": reports}, f, ensure_ascii=False, indent=1)

print("권역 %d개 → %s" % (len(reports), os.path.abspath(OUT)))
for k, v in reports.items():
    print("  - %s : 유동avg %s · 지하철 %s · 버스정류장 %s · 매체 %d · 연령%s" % (
        k, v.get("footAvg"), v.get("subwayAvg") or (v.get("summary") and "요약"),
        len(v.get("busStops", [])), len(v.get("media", [])), "O" if v.get("age") else "X"))
